import os
import subprocess
from .base_converter import BaseConverter

try:
    import pypandoc
    PYPANDOC_AVAILABLE = True
except ImportError:
    PYPANDOC_AVAILABLE = False

try:
    from docx import Document
    from docx.shared import Inches
    PYTHON_DOCX_AVAILABLE = True
except ImportError:
    PYTHON_DOCX_AVAILABLE = False

try:
    import openpyxl
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

class DocumentConverter(BaseConverter):
    def __init__(self):
        super().__init__()
        self.supported_input_formats = [
            'docx', 'doc', 'pdf', 'txt', 'odt', 'rtf', 'html', 'md',
            'xlsx', 'xls', 'csv', 'ods',
            'pptx', 'ppt', 'odp'
        ]
        self.supported_output_formats = [
            'pdf', 'docx', 'txt', 'rtf', 'html', 'odt',
            'xlsx', 'csv', 'ods'
        ]

    def convert(self, input_path, output_format, output_directory):
        """Convert document files using appropriate method."""
        try:
            if not self.ensure_output_directory(output_directory):
                return False

            input_format = self.get_file_extension(input_path)
            output_path = self.get_output_filename(input_path, output_format, output_directory)

            # Use pypandoc for most document conversions
            if PYPANDOC_AVAILABLE and self._can_use_pandoc(input_format, output_format):
                return self._convert_with_pandoc(input_path, output_path, output_format)

            # Fallback methods for specific conversions
            if input_format in ['xlsx', 'xls', 'csv'] and output_format == 'csv':
                return self._convert_spreadsheet_to_csv(input_path, output_path)

            if input_format == 'txt' and output_format == 'docx':
                return self._convert_txt_to_docx(input_path, output_path)

            # Try LibreOffice conversion as fallback
            return self._convert_with_libreoffice(input_path, output_format, output_directory)

        except Exception as e:
            print(f"Document conversion error: {str(e)}")
            return False

    def _can_use_pandoc(self, input_format, output_format):
        """Check if pandoc can handle this conversion."""
        pandoc_formats = ['docx', 'pdf', 'txt', 'html', 'odt', 'rtf', 'md']
        return input_format in pandoc_formats and output_format in pandoc_formats

    def _convert_with_pandoc(self, input_path, output_path, output_format):
        """Convert using pypandoc."""
        try:
            pypandoc.convert_file(input_path, output_format, outputfile=output_path)
            return os.path.exists(output_path)
        except Exception as e:
            print(f"Pandoc conversion error: {str(e)}")
            return False

    def _convert_with_libreoffice(self, input_path, output_format, output_directory):
        """Convert using LibreOffice headless mode."""
        try:
            # Map output formats to LibreOffice filter names
            lo_filters = {
                'pdf': 'writer_pdf_Export',
                'docx': 'MS Word 2007 XML',
                'odt': 'writer8',
                'rtf': 'Rich Text Format',
                'txt': 'Text'
            }

            if output_format not in lo_filters:
                return False

            cmd = [
                'libreoffice',
                '--headless',
                '--convert-to', output_format,
                '--outdir', output_directory,
                input_path
            ]

            result = subprocess.run(cmd, capture_output=True, text=True, timeout=60)
            return result.returncode == 0

        except Exception as e:
            print(f"LibreOffice conversion error: {str(e)}")
            return False

    def _convert_spreadsheet_to_csv(self, input_path, output_path):
        """Convert spreadsheet to CSV."""
        if not OPENPYXL_AVAILABLE:
            return False

        try:
            if input_path.endswith('.csv'):
                # Just copy the file
                import shutil
                shutil.copy2(input_path, output_path)
                return True

            # Convert Excel to CSV
            wb = openpyxl.load_workbook(input_path)
            ws = wb.active

            import csv
            with open(output_path, 'w', newline='', encoding='utf-8') as csvfile:
                writer = csv.writer(csvfile)
                for row in ws.iter_rows(values_only=True):
                    writer.writerow(row)

            return True

        except Exception as e:
            print(f"Spreadsheet conversion error: {str(e)}")
            return False

    def _convert_txt_to_docx(self, input_path, output_path):
        """Convert text file to DOCX."""
        if not PYTHON_DOCX_AVAILABLE:
            return False

        try:
            doc = Document()

            with open(input_path, 'r', encoding='utf-8') as f:
                content = f.read()

            # Split content into paragraphs
            paragraphs = content.split('\n\n')
            for paragraph in paragraphs:
                if paragraph.strip():
                    doc.add_paragraph(paragraph.strip())

            doc.save(output_path)
            return True

        except Exception as e:
            print(f"TXT to DOCX conversion error: {str(e)}")
            return False